"""Golden compliance case: PS-100 spec requirements vs the CX-250 supplier data sheet."""

import io

from eb.compliance.models import LLMVerdict
from eb.jobs import run_pending_once


def _upload(client, text: str, filename: str, **form):
    r = client.post(
        "/documents", files={"file": (filename, io.BytesIO(text.encode()), "text/markdown")}, data=form
    )
    assert r.status_code == 201, r.text
    return r.json()


def _ingest_both(client, sample_spec, sample_datasheet):
    spec = _upload(client, sample_spec, "ps-100.md", doc_number="PS-100", revision="A", doc_type="spec")
    ds = _upload(client, sample_datasheet, "cx-250.md", doc_number="DS-CX-250", doc_type="datasheet")
    assert run_pending_once() >= 2
    return spec, ds


def _verdict(verdict, quote, confidence=0.9, values=None, rationale="scripted"):
    return LLMVerdict(
        verdict=verdict,
        confidence=confidence,
        quotes=[{"source_index": 0, "quote": quote}] if quote else [],
        found_values=[{"param": p, "value_text": v, "source_index": 0} for p, v in (values or [])],
        rationale=rationale,
    )


def test_extraction_is_deterministic_without_llm(client, sample_spec, sample_datasheet, fake_llm):
    spec, _ = _ingest_both(client, sample_spec, sample_datasheet)
    s = client.post("/compliance/sets", json={"name": "PS-100 requirements", "discipline": ["mechanical"]})
    assert s.status_code == 201, s.text
    set_id = s.json()["id"]

    r = client.post(f"/compliance/sets/{set_id}/extract", json={"document_id": spec["document_id"]})
    assert r.status_code == 200, r.text
    assert r.json()["added"] == 5
    # FakeLLM.parse raised (nothing scripted); extraction fell back to the deterministic path.
    assert any(c.get("schema") is not None for c in fake_llm.calls)

    body = client.get(f"/compliance/sets/{set_id}").json()
    assert body["set"]["source"] == "extracted"
    reqs = {q["req_id"]: q for q in body["requirements"]}
    assert list(reqs) == ["4.1", "4.2", "4.3", "4.4", "5.1"]
    assert reqs["4.1"]["text"].startswith("The pump shall deliver")
    q41 = reqs["4.1"]["quantity"]
    assert q41["op"] == ">=" and q41["value"] == 250 and q41["unit"] == "m3/h"
    assert q41["and"][0]["value"] == 45 and q41["and"][0]["unit"] == "m"
    q42 = reqs["4.2"]["quantity"]
    assert (q42["op"], q42["value"], q42["unit"]) == (">=", 16, "bar(g)")
    q43 = reqs["4.3"]["quantity"]
    assert q43["value"] == 400 and q43["unit"] == "V"
    assert [x["value"] for x in q43["and"]] == [50, 1.15]
    assert reqs["4.4"]["quantity"] is None and reqs["5.1"]["quantity"] is None
    assert reqs["4.2"]["ears_pattern"] == "ubiquitous"
    assert reqs["4.1"]["source_chunk_id"]

    sets = client.get("/compliance/sets").json()
    assert sets[0]["id"] == set_id and sets[0]["requirement_count"] == 5

    assert client.delete(f"/compliance/sets/{set_id}").status_code == 204
    assert client.get(f"/compliance/sets/{set_id}").status_code == 404


def test_golden_run_verifies_quotes_and_overrides_numbers(client, sample_spec, sample_datasheet, fake_llm):
    spec, ds = _ingest_both(client, sample_spec, sample_datasheet)
    set_id = client.post("/compliance/sets", json={"name": "PS-100"}).json()["id"]
    assert (
        client.post(f"/compliance/sets/{set_id}/extract", json={"document_id": spec["document_id"]}).json()[
            "added"
        ]
        == 5
    )
    # A manual text requirement the data sheet says nothing about: the scripted LLM will fabricate a quote.
    r = client.post(
        f"/compliance/sets/{set_id}/requirements",
        json={
            "requirements": [
                {"req_id": "6.1", "text": "The pump shall be painted RAL 5015 blue.", "kind": "design"},
            ]
        },
    )
    assert r.json() == {"added": 1}

    # Scripted verdicts in requirement order (4.1, 4.2, 4.3, 4.4, 5.1, 6.1); quotes verbatim from the
    # data sheet.
    fake_llm.parsed.extend(
        [
            _verdict(
                "compliant",
                "Rated capacity: 250 m3/h at 48 m head.",
                values=[("rated flow", "250 m3/h"), ("head", "48 m")],
            ),
            _verdict(
                "compliant",
                "Maximum allowable working pressure of casing: 20 bar(g) at 38 degC.",
                values=[("design pressure", "20 bar(g)")],
            ),
            # The LLM wrongly passes the service factor; code must override.
            _verdict("compliant", "Motor: 400 V, 50 Hz, three phase, service factor 1.0.", confidence=0.95),
            _verdict(
                "non_compliant",
                "Casing material: ductile iron ASTM A395.",
                confidence=0.92,
                rationale="Ductile iron A395 is not ASTM A216 WCB carbon steel.",
            ),
            _verdict(
                "partial",
                "Performance test: ISO 9906 Grade 2B, witnessed on request.",
                confidence=0.7,
                rationale="Witnessing is offered on request, not committed.",
            ),
            _verdict("compliant", "Paint finish: RAL 5015 blue epoxy.", confidence=0.9),  # not in the source
        ]
    )

    r = client.post("/compliance/runs", json={"set_id": set_id, "document_id": ds["document_id"]})
    assert r.status_code == 201, r.text
    run_id = r.json()["run_id"]
    assert r.json()["job_id"]
    assert run_pending_once() == 1
    assert fake_llm.parsed == []

    body = client.get(f"/compliance/runs/{run_id}").json()
    assert body["run"]["status"] == "done", body["run"]
    by_id = {f["requirement"]["req_id"]: f for f in body["findings"]}
    assert list(by_id) == ["4.1", "4.2", "4.3", "4.4", "5.1", "6.1"]

    f41 = by_id["4.1"]
    assert f41["verdict"] == "compliant" and f41["llm_verdict"] == "compliant"
    assert f41["numeric_check"]["result"] == "pass"
    assert len(f41["numeric_check"]["checks"]) == 2
    assert f41["evidence"][0]["verified"] is True and f41["evidence"][0]["page"] == 1
    assert f41["needs_review"] is False

    f42 = by_id["4.2"]
    assert f42["verdict"] == "compliant" and f42["numeric_check"]["result"] == "pass"
    assert "20 bar(g)" in f42["numeric_check"]["detail"]

    f43 = by_id["4.3"]
    assert f43["llm_verdict"] == "compliant"
    assert f43["verdict"] == "non_compliant"  # code overrides the LLM on numeric clauses
    assert f43["numeric_check"]["result"] == "fail" and f43["numeric_check"]["overrode_llm"] is True
    assert "1.15" in f43["numeric_check"]["detail"]
    assert f43["needs_review"] is True  # LLM and code disagree -> a human looks at it

    f44 = by_id["4.4"]
    assert f44["verdict"] == "non_compliant" and f44["numeric_check"] is None
    assert f44["evidence"][0]["quote"] == "Casing material: ductile iron ASTM A395."
    assert f44["needs_review"] is False

    f51 = by_id["5.1"]
    assert f51["verdict"] in ("partial", "insufficient_evidence") and f51["needs_review"] is True

    f61 = by_id["6.1"]
    assert f61["verdict"] == "unverified" and f61["llm_verdict"] == "compliant"
    assert f61["evidence"] == [] and f61["needs_review"] is True
    assert "dropped" in f61["rationale"]

    summary = body["summary"]
    assert summary["false_compliant_guard"] == 0
    assert summary["total"] == 6 and summary["needs_review"] == 3
    assert summary["verdicts"] == {"compliant": 2, "non_compliant": 2, "partial": 1, "unverified": 1}
    assert summary["unrequested_claims"] == []  # the data sheet has no "shall" chunks

    runs = client.get("/compliance/runs", params={"set_id": set_id}).json()
    assert [x["id"] for x in runs] == [run_id] and runs[0]["finding_count"] == 6

    # Reviewer override.
    r = client.post(
        f"/compliance/findings/{f61['id']}/review",
        json={"verdict": "insufficient_evidence", "note": "No paint data in the data sheet."},
    )
    assert r.status_code == 200, r.text
    assert r.json()["reviewer_verdict"] == "insufficient_evidence" and r.json()["needs_review"] is False
    assert r.json()["verdict"] == "unverified"  # the machine verdict is kept

    # Excel export.
    r = client.get(f"/compliance/runs/{run_id}/export.xlsx")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    assert "spreadsheetml" in r.headers["content-type"]
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Compliance matrix"]
    rows = list(ws.iter_rows(values_only=True))
    header = next(r for r in rows if r and r[0] == "Req ID")
    assert list(header) == [
        "Req ID",
        "Requirement",
        "Verdict",
        "Reviewer verdict",
        "Confidence",
        "Evidence quotes",
        "Pages",
        "Numeric check",
        "Rationale",
        "Needs review",
    ]
    data = {r[0]: r for r in rows[rows.index(header) + 1 :]}
    assert data["4.3"][2] == "non_compliant" and data["6.1"][3] == "insufficient_evidence"


def test_run_without_llm_routes_everything_to_review(client, sample_spec, sample_datasheet, fake_llm):
    spec, ds = _ingest_both(client, sample_spec, sample_datasheet)
    set_id = client.post("/compliance/sets", json={"name": "PS-100"}).json()["id"]
    client.post(f"/compliance/sets/{set_id}/extract", json={"document_id": spec["document_id"]})
    run_id = client.post(
        "/compliance/runs", json={"set_id": set_id, "document_id": ds["document_id"]}
    ).json()["run_id"]
    run_pending_once()
    body = client.get(f"/compliance/runs/{run_id}").json()
    assert body["run"]["status"] == "done"
    assert all(f["verdict"] == "insufficient_evidence" and f["needs_review"] for f in body["findings"])
    assert body["summary"]["llm_errors"] == 5 and body["summary"]["false_compliant_guard"] == 0


def test_import_csv_and_xlsx(client, fake_llm):
    set_id = client.post("/compliance/sets", json={"name": "imported"}).json()["id"]
    csv_text = (
        "ID,Requirement,Kind,Param,Op,Value,Unit,Tol\n"
        "R-1,The pump shall deliver not less than 250 m3/h.,performance,rated flow,>=,250,m3/h,\n"
        "R-2,The casing shall be carbon steel.,constraint,,,,,\n"
    )
    r = client.post(
        f"/compliance/sets/{set_id}/import",
        files={"file": ("reqs.csv", io.BytesIO(csv_text.encode()), "text/csv")},
    )
    assert r.status_code == 200, r.text
    assert r.json() == {"added": 2}

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["req_id", "text", "param", "op", "value", "unit"])
    ws.append(["R-3", "Noise shall not exceed 85 dB(A).", "noise", "<=", 85, "dB(A)"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(
        f"/compliance/sets/{set_id}/import",
        files={
            "file": (
                "reqs.xlsx",
                io.BytesIO(buf.getvalue()),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.json() == {"added": 1}

    body = client.get(f"/compliance/sets/{set_id}").json()
    reqs = {q["req_id"]: q for q in body["requirements"]}
    assert set(reqs) == {"R-1", "R-2", "R-3"}
    assert reqs["R-1"]["quantity"]["op"] == ">=" and reqs["R-1"]["quantity"]["value"] == 250
    assert reqs["R-2"]["quantity"] is None and reqs["R-2"]["kind"] == "constraint"
    assert reqs["R-3"]["quantity"]["unit"] == "dB(A)"
    assert body["set"]["source"] == "excel"


def test_acl_blocks_runs_on_hidden_documents(client, sample_spec, sample_datasheet, db, fake_llm):
    spec, ds = _ingest_both(client, sample_spec, sample_datasheet)
    set_id = client.post("/compliance/sets", json={"name": "PS-100"}).json()["id"]
    with db() as conn:
        conn.execute("DELETE FROM document_acl WHERE document_id = %s", (ds["document_id"],))
    r = client.post("/compliance/runs", json={"set_id": set_id, "document_id": ds["document_id"]})
    assert r.status_code == 404
    with db() as conn:
        conn.execute("DELETE FROM document_acl WHERE document_id = %s", (spec["document_id"],))
    r = client.post(f"/compliance/sets/{set_id}/extract", json={"document_id": spec["document_id"]})
    assert r.status_code == 404
