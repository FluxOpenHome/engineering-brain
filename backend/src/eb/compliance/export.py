"""Compliance matrix export (Excel)."""

from __future__ import annotations

import io
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from eb.db import connection

COLUMNS = [
    "Req ID",
    "Requirement",
    "Verdict",
    "Reviewer verdict",
    "Confidence",
    "Evidence quotes",
    "Pages",
    "Numeric check",
    "Rationale",
    "Needs review",
]
_FILL = {
    "compliant": "C6EFCE",
    "non_compliant": "FFC7CE",
    "partial": "FFEB9C",
    "insufficient_evidence": "D9D9D9",
    "unverified": "F4B084",
    "not_applicable": "EDEDED",
}


def load_matrix(run_id: UUID) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    with connection() as conn:
        run = conn.execute(
            """
            SELECT r.*, s.name AS set_name, d.doc_number, d.title AS document_title, rev.revision
            FROM compliance_runs r
            JOIN requirement_sets s ON s.id = r.set_id
            JOIN documents d ON d.id = r.subject_document_id
            JOIN document_revisions rev ON rev.id = r.subject_revision_id
            WHERE r.id = %s
            """,
            (run_id,),
        ).fetchone()
        if run is None:
            raise KeyError("run not found")
        rows = conn.execute(
            """
            SELECT f.*, q.req_id, q.text AS requirement_text, q.kind::text AS requirement_kind, q.quantity
            FROM compliance_findings f JOIN requirements q ON q.id = f.requirement_id
            WHERE f.run_id = %s ORDER BY q.ordinal, q.created_at, q.req_id
            """,
            (run_id,),
        ).fetchall()
    return run, rows


def export_xlsx(run_id: UUID) -> bytes:
    run, rows = load_matrix(run_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance matrix"
    ws.append(
        [
            f"Compliance matrix: {run['set_name']} vs {run['doc_number'] or ''} {run['document_title']} "
            f"Rev {run['revision']}"
        ]
    )
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"Run {run['id']} - status {run['status']} - model {run['model'] or ''}"])
    ws.append([])
    ws.append(COLUMNS)
    head_row = ws.max_row
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=head_row, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in rows:
        ev = r["evidence"] or []
        quotes = "\n".join(f"“{e.get('quote', '')}”" for e in ev if e.get("verified"))
        pages = ", ".join(
            sorted({str(e["page"]) for e in ev if e.get("page") is not None}, key=lambda x: int(x))
        )
        nc = r["numeric_check"]
        nc_text = f"{nc['result']}: {nc['detail']}" if nc else ""
        ws.append(
            [
                r["req_id"],
                r["requirement_text"],
                r["verdict"],
                r["reviewer_verdict"] or "",
                round(r["confidence"], 2) if r["confidence"] is not None else "",
                quotes,
                pages,
                nc_text,
                r["rationale"],
                "yes" if r["needs_review"] else "no",
            ]
        )
        fill = _FILL.get(r["reviewer_verdict"] or r["verdict"])
        if fill:
            ws.cell(row=ws.max_row, column=3).fill = PatternFill("solid", fgColor=fill)
    widths = [12, 60, 18, 18, 11, 60, 8, 50, 60, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=head_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    sm = wb.create_sheet("Summary")
    summary = run["summary"] or {}
    sm.append(["Metric", "Value"])
    sm["A1"].font = sm["B1"].font = Font(bold=True)
    for key in (
        "total",
        "needs_review",
        "false_compliant_guard",
        "llm_errors",
        "evidence_chunks_cited",
        "unrequested_claims_count",
    ):
        if key in summary:
            sm.append([key, summary[key]])
    for verdict, n in (summary.get("verdicts") or {}).items():
        sm.append([f"verdict: {verdict}", n])
    sm.append([])
    sm.append(["Unrequested claims (subject 'shall' statements not cited by any finding)"])
    sm.append(["Chunk", "Clause", "Page", "Text"])
    for c in summary.get("unrequested_claims") or []:
        sm.append([c.get("chunk_id"), c.get("clause_id"), c.get("page"), c.get("text")])
    sm.column_dimensions["A"].width = 40
    sm.column_dimensions["D"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
