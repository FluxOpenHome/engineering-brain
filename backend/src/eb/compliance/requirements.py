"""Requirement sets and requirements: CRUD, table import and extraction from a document."""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any
from uuid import UUID

import structlog

from eb.auth import Principal
from eb.compliance import units
from eb.compliance.models import ExtractedRequirements, QuantitySpec, RequirementIn
from eb.db import connection
from eb.llm import DocumentBlock, get_llm
from eb.specs.rules import classify_ears

log = structlog.get_logger(__name__)

KINDS = {"functional", "performance", "interface", "constraint", "design", "other"}

# ------------------------------------------------------------------------- sets


def create_set(
    name: str,
    *,
    description: str = "",
    discipline: list[str] | None = None,
    user: Principal | None = None,
    source: str = "manual",
    document_id: UUID | None = None,
) -> dict[str, Any]:
    with connection() as conn:
        return conn.execute(
            """
            INSERT INTO requirement_sets (name, description, discipline, source, document_id, created_by)
            VALUES (%s, %s, %s, %s, %s, %s) RETURNING *
            """,
            (name, description, discipline or [], source, document_id, user.id if user else None),
        ).fetchone()


def list_sets() -> list[dict[str, Any]]:
    with connection() as conn:
        return conn.execute(
            """
            SELECT s.*, (SELECT count(*) FROM requirements r WHERE r.set_id = s.id) AS requirement_count,
                   d.doc_number AS document_number, d.title AS document_title
            FROM requirement_sets s LEFT JOIN documents d ON d.id = s.document_id
            ORDER BY s.created_at DESC
            """
        ).fetchall()


def get_set(set_id: UUID) -> dict[str, Any] | None:
    with connection() as conn:
        s = conn.execute(
            """
            SELECT s.*, d.doc_number AS document_number, d.title AS document_title
            FROM requirement_sets s LEFT JOIN documents d ON d.id = s.document_id WHERE s.id = %s
            """,
            (set_id,),
        ).fetchone()
        if s is None:
            return None
        reqs = conn.execute(
            "SELECT * FROM requirements WHERE set_id = %s ORDER BY ordinal, created_at, req_id", (set_id,)
        ).fetchall()
    return {"set": s, "requirements": reqs}


def delete_set(set_id: UUID) -> bool:
    with connection() as conn:
        row = conn.execute("DELETE FROM requirement_sets WHERE id = %s RETURNING id", (set_id,)).fetchone()
    return row is not None


# ------------------------------------------------------------------------- requirements


def _quantity_json(q: QuantitySpec | dict[str, Any] | None) -> str | None:
    if q is None:
        return None
    if isinstance(q, QuantitySpec):
        return json.dumps(q.to_json())
    return json.dumps(q)


def _next_ordinal(conn, set_id: UUID) -> int:
    row = conn.execute(
        "SELECT coalesce(max(ordinal), -1) + 1 AS n FROM requirements WHERE set_id = %s", (set_id,)
    ).fetchone()
    return int(row["n"])


def _unique_req_id(conn, set_id: UUID, req_id: str, taken: set[str]) -> str:
    base = req_id.strip() or "REQ"
    cand = base
    i = 2
    while (
        cand in taken
        or conn.execute(
            "SELECT 1 FROM requirements WHERE set_id = %s AND req_id = %s", (set_id, cand)
        ).fetchone()
    ):
        cand = f"{base} ({i})"
        i += 1
    taken.add(cand)
    return cand


def add_requirements(
    set_id: UUID,
    items: list[RequirementIn | dict[str, Any]],
    *,
    source_chunk_ids: list[UUID | None] | None = None,
) -> int:
    """Insert requirements; req_ids that collide inside the set get a numeric suffix."""
    added = 0
    taken: set[str] = set()
    with connection() as conn:
        if conn.execute("SELECT 1 FROM requirement_sets WHERE id = %s", (set_id,)).fetchone() is None:
            raise KeyError("requirement set not found")
        ordinal = _next_ordinal(conn, set_id)
        for i, item in enumerate(items):
            r = item if isinstance(item, RequirementIn) else RequirementIn.model_validate(item)
            kind = r.kind if r.kind in KINDS else _guess_kind(r.text, r.quantity is not None)
            req_id = _unique_req_id(conn, set_id, r.req_id, taken)
            chunk_id = source_chunk_ids[i] if source_chunk_ids and i < len(source_chunk_ids) else None
            conn.execute(
                """
                INSERT INTO requirements (set_id, req_id, text, kind, quantity, source_chunk_id, ears_pattern,
                                          ordinal)
                VALUES (%s, %s, %s, %s, %s::jsonb, %s, %s, %s)
                """,
                (
                    set_id,
                    req_id,
                    r.text.strip(),
                    kind,
                    _quantity_json(r.quantity),
                    chunk_id,
                    classify_ears(r.text),
                    ordinal,
                ),
            )
            ordinal += 1
            added += 1
    return added


def _guess_kind(text: str, has_quantity: bool) -> str:
    low = text.lower()
    if has_quantity:
        return "performance"
    if re.search(r"\b(interface|connect|connection|supply from|terminal|signal|tie-in)\b", low):
        return "interface"
    if re.search(r"\b(material|standard|accordance|comply|complian|code|certif|iso|iec|astm|api)\b", low):
        return "constraint"
    if re.search(r"\b(design|dimension|arrangement|layout|mount|paint|coating|colour|color)\b", low):
        return "design"
    return "functional"


# ------------------------------------------------------------------------- table import

_COLS = {
    "req_id": {"id", "req_id", "req id", "requirement id", "reqid", "ref", "clause", "no", "no."},
    "text": {"text", "requirement", "requirement text", "description", "statement"},
    "kind": {"kind", "type", "category"},
    "param": {"param", "parameter"},
    "op": {"op", "operator", "comparison"},
    "value": {"value", "val", "limit"},
    "unit": {"unit", "units", "uom"},
    "tol": {"tol", "tolerance"},
    "high": {"high", "max", "upper", "value_high"},
}


def _rows_from_bytes(data: bytes, filename: str) -> list[list[Any]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect=dialect)]


def parse_table(data: bytes, filename: str) -> list[RequirementIn]:
    rows = _rows_from_bytes(data, filename)
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        return []
    header = [str(c or "").strip().lower() for c in rows[0]]
    idx: dict[str, int] = {}
    for key, names in _COLS.items():
        for i, h in enumerate(header):
            if h in names and key not in idx:
                idx[key] = i
    if "text" not in idx:
        raise ValueError("table needs a 'text' or 'requirement' column")

    def cell(row: list[Any], key: str) -> str | None:
        i = idx.get(key)
        if i is None or i >= len(row) or row[i] in (None, ""):
            return None
        return str(row[i]).strip()

    out: list[RequirementIn] = []
    for n, row in enumerate(rows[1:], start=1):
        text = cell(row, "text")
        if not text:
            continue
        req_id = cell(row, "req_id") or f"R-{n}"
        kind = (cell(row, "kind") or "").lower() or None
        if kind and kind not in KINDS:
            kind = None
        quantity: QuantitySpec | None = None
        value = cell(row, "value")
        if value:
            try:
                qdict: dict[str, Any] = {
                    "param": cell(row, "param"),
                    "op": _norm_op(cell(row, "op")),
                    "value": float(value.replace(",", "")),
                    "unit": cell(row, "unit"),
                    "tol": cell(row, "tol"),
                }
                high = cell(row, "high")
                if high:
                    qdict["high"] = float(high.replace(",", ""))
                    if qdict["op"] == "==":
                        qdict["op"] = "between"
                quantity = QuantitySpec.model_validate(qdict)
            except (ValueError, TypeError):
                quantity = None
        out.append(RequirementIn(req_id=req_id, text=text, kind=kind, quantity=quantity))
    return out


def _norm_op(op: str | None) -> str:
    if not op:
        return "=="
    o = op.strip().lower()
    table = {
        ">=": ">=",
        "≥": ">=",
        "=>": ">=",
        "min": ">=",
        "minimum": ">=",
        "not less than": ">=",
        "at least": ">=",
        "<=": "<=",
        "≤": "<=",
        "=<": "<=",
        "max": "<=",
        "maximum": "<=",
        "not more than": "<=",
        "at most": "<=",
        "==": "==",
        "=": "==",
        "eq": "==",
        "equal": "==",
        ">": ">",
        "gt": ">",
        "<": "<",
        "lt": "<",
        "between": "between",
        "range": "between",
    }
    return table.get(o, "==")


def import_table(set_id: UUID, file_bytes: bytes, filename: str) -> int:
    items = parse_table(file_bytes, filename)
    n = add_requirements(set_id, items) if items else 0
    with connection() as conn:
        conn.execute(
            "UPDATE requirement_sets SET source = 'excel' WHERE id = %s AND source = 'manual' "
            "AND NOT EXISTS (SELECT 1 FROM requirements r WHERE r.set_id = %s "
            "AND r.source_chunk_id IS NOT NULL)",
            (set_id, set_id),
        )
    return n


# ------------------------------------------------------------------------- extraction from a document

_GE_WORDS = re.compile(
    r"\b(not less than|no less than|at least|minimum|min\.|or more|or greater|or higher|"
    r"shall deliver|deliver|"
    r"capable of|capacity of|withstand|exceed(?:s|ing)? at least)\b|≥|>=",
    re.IGNORECASE,
)
_LE_WORDS = re.compile(
    r"\b(not more than|no more than|not exceed(?:ing)?|shall not exceed|at most|maximum|max\.|"
    r"or less|or lower|"
    r"up to|limited to|within)\b|≤|<=",
    re.IGNORECASE,
)
_CLAUSE_PREFIX_RE = re.compile(r"^\s*((?:[A-Z]\.)?\d+(?:\.\d+){1,4})\s+")
_STOP_PARAM = {
    "a",
    "an",
    "the",
    "of",
    "at",
    "to",
    "for",
    "with",
    "and",
    "or",
    "be",
    "shall",
    "must",
    "is",
    "are",
    "its",
    "not",
    "less",
    "more",
    "than",
    "least",
    "minimum",
    "maximum",
    "min",
    "max",
    "rated",
    "have",
    "has",
    "by",
}


def _param_name(text: str, q: units.Quantity, n_quantities: int) -> str | None:
    """Noun phrase naming the parameter: the subject before 'shall' for single-value clauses, else
    the few content words preceding the number."""
    if n_quantities == 1:
        m = re.match(r"^\s*(?:the\s+)?(.+?)\s+(?:shall|must)\b", text, re.IGNORECASE)
        if m and len(m.group(1).split()) <= 8:
            return m.group(1).strip().rstrip(",")
    words = [
        w
        for w in re.findall(r"[A-Za-z][A-Za-z\-]*", q.context)
        if w.lower() not in _STOP_PARAM
        and len(w) > 1
        and not units.is_unit_token(w)
        and not _GE_WORDS.fullmatch(w)
        and not _LE_WORDS.fullmatch(w)
    ]
    return " ".join(words[-3:]) or None


def deterministic_quantity(text: str) -> QuantitySpec | None:
    """Fill `quantity` from the clause text without an LLM.

    The operator comes from the number's own prefix ("not less than 16 bar"), else from a
    clause-level keyword (minimum/at least/deliver -> ">=", maximum/not exceed -> "<="), else "==".
    Clauses with several values carry the extra ones under "and".
    """
    qs = [
        q for q in units.parse_quantity(text) if not (q.unit is None and "." not in q.text and q.op is None)
    ]
    if not qs or len(qs) > 4:
        return None
    clause_op = ">=" if _GE_WORDS.search(text) else "<=" if _LE_WORDS.search(text) else "=="
    specs: list[QuantitySpec] = []
    for q in qs:
        op = q.op or clause_op
        tol: float | str | None = None
        if q.tol_rel is not None:
            tol = f"{q.tol_rel * 100:g}%"
        elif q.tol is not None:
            tol = q.tol
        unit_text = q.unit_text
        if (
            q.gauge
            and unit_text
            and not unit_text.endswith(f"({q.gauge})")
            and not unit_text.endswith(q.gauge)
        ):
            unit_text = f"{unit_text}({q.gauge})"
        specs.append(
            QuantitySpec(
                param=_param_name(text, q, len(qs)),
                op=op,
                value=q.value,
                unit=unit_text,
                tol=tol,
                high=q.high if op == "between" else None,
            )
        )
    primary = specs[0]
    primary.and_ = specs[1:]
    return primary


EXTRACT_SYSTEM = """You extract atomic engineering requirements from specification clauses.

Rules:
- Each <source> is one clause. Produce one requirement per distinct obligation; split clauses that \
contain several obligations (flow AND head -> two requirements) and set source_index to the clause.
- Keep the wording of the source; do not paraphrase numbers, units, standards or clause ids.
- req_id: the clause id from the source title, with a letter suffix (4.1a, 4.1b) when split.
- kind: functional | performance | interface | constraint | design | other.
- quantity: when the requirement fixes a number, give {param, op, value, unit, tol}. op is one of \
>=, <=, ==, >, <, between ("not less than" -> ">=", "maximum" -> "<=", a nominal value -> "=="). \
Write the unit as in the source (bar(g), m3/h, degC). Leave quantity null for text-only requirements.
- ears_pattern: ubiquitous | event-driven | state-driven | optional | unwanted | complex.
- Never invent requirements that are not in the sources."""


def _clause_chunks(conn, revision_id: UUID) -> list[dict[str, Any]]:
    return conn.execute(
        """
        SELECT id, ordinal, kind::text AS kind, clause_id, text, page_start, section_path
        FROM chunks WHERE revision_id = %s AND kind IN ('requirement', 'clause')
        ORDER BY ordinal
        """,
        (revision_id,),
    ).fetchall()


def _strip_clause_prefix(text: str, clause_id: str | None) -> str:
    if clause_id:
        m = _CLAUSE_PREFIX_RE.match(text)
        if m and m.group(1) == clause_id:
            return text[m.end() :].strip()
    return text.strip()


def _llm_extract(chunks: list[dict[str, Any]]) -> ExtractedRequirements | None:
    """Ask the LLM to atomise clauses; None when no LLM output is available."""
    docs = [DocumentBlock(text=c["text"], title=f"clause {c['clause_id'] or c['ordinal']}") for c in chunks]
    try:
        return get_llm().parse(
            schema=ExtractedRequirements,
            system=EXTRACT_SYSTEM,
            user="Extract the atomic requirements from every clause above. Return them in source order.",
            documents=docs,
        )
    except RuntimeError as e:  # FakeLLM without a script, or a provider that refused
        log.info("compliance.extract.llm_unavailable", error=str(e))
        return None
    except Exception as e:  # pragma: no cover - provider errors fall back to deterministic extraction
        log.warning("compliance.extract.llm_failed", error=str(e))
        return None


def extract_from_document(set_id: UUID, document_id: UUID, user: Principal, *, use_llm: bool = True) -> int:
    """Create requirements from the active revision's requirement/clause chunks.

    Deterministic: one requirement per chunk, req_id = clause_id. When the LLM is available it
    atomises multi-requirement clauses and fills `quantity`; otherwise `quantity` is filled by
    `deterministic_quantity`.
    """
    with connection() as conn:
        doc = conn.execute(
            """
            SELECT d.id, d.current_revision_id FROM documents d
            WHERE d.id = %s AND EXISTS (SELECT 1 FROM document_acl a WHERE a.document_id = d.id
                                        AND a.group_id = ANY(%s))
            """,
            (document_id, user.group_ids),
        ).fetchone()
        if doc is None:
            raise KeyError("document not found")
        if doc["current_revision_id"] is None:
            raise ValueError("document has no indexed revision yet")
        revision_id = doc["current_revision_id"]
        chunks = _clause_chunks(conn, revision_id)
    if not chunks:
        return 0

    by_source: dict[int, list[Any]] = {}
    llm_out = _llm_extract(chunks) if use_llm else None
    if llm_out is not None:
        for item in llm_out.requirements:
            if item.source_index is None or not (0 <= item.source_index < len(chunks)):
                continue
            by_source.setdefault(item.source_index, []).append(item)

    items: list[RequirementIn] = []
    chunk_ids: list[UUID | None] = []
    for i, c in enumerate(chunks):
        fallback_id = c["clause_id"] or f"C-{c['ordinal'] + 1}"
        text = _strip_clause_prefix(c["text"], c["clause_id"])
        if by_source.get(i):
            for item in by_source[i]:
                items.append(
                    RequirementIn(
                        req_id=item.req_id or fallback_id,
                        text=item.text or text,
                        kind=item.kind,
                        quantity=item.quantity,
                    )
                )
                chunk_ids.append(c["id"])
            continue
        q = deterministic_quantity(text)
        items.append(
            RequirementIn(req_id=fallback_id, text=text, kind=_guess_kind(text, q is not None), quantity=q)
        )
        chunk_ids.append(c["id"])

    added = add_requirements(set_id, items, source_chunk_ids=chunk_ids)
    with connection() as conn:
        conn.execute(
            """
            UPDATE requirement_sets
            SET source = CASE WHEN source = 'manual' THEN 'extracted' ELSE source END,
                document_id = coalesce(document_id, %s), revision_id = coalesce(revision_id, %s)
            WHERE id = %s
            """,
            (document_id, revision_id, set_id),
        )
    return added
